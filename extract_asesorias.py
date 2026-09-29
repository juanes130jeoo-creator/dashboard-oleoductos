import openpyxl
import json
import datetime
import sys
import re

def get_str(val):
    if val is None: return ""
    return str(val).strip()

def get_date(val):
    if val is None: return None
    if isinstance(val, datetime.datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, str):
        try:
            return datetime.datetime.strptime(val, "%Y-%m-%d %H:%M:%S").strftime("%Y-%m-%d")
        except:
            return val.split(" ")[0]
    return str(val)

def normalize_spaces(text):
    if not text: return ""
    return re.sub(r'\s+', ' ', text).strip()

CATALOGO = {
    "comercial": "Comercial",
    "comercial/ innovación": "Comercial / Innovación",
    "contabilidad": "Contabilidad",
    "finanzas": "Finanzas",
    "marketing": "Marketing",
    "marketing digital": "Marketing Digital",
    "modelo canvas": "Modelo Canvas",
    "posicionamiento": "Posicionamiento",
    "redes sociales": "Redes Sociales",
    "redes sociales - tik tok": "Redes Sociales",
    "redes sociales/ portafolio": "Redes Sociales",
    "portafolio/redes sociales": "Redes Sociales",
    "servicio al cliente": "Servicio al Cliente",
    "servicio la cliente": "Servicio al Cliente",
    "transformación digital": "Transformación Digital",
    "marketing o contabilidad": "Sin Definir / Pendiente",
    "": "Sin Definir / Pendiente"
}

def normalize_tema(tema):
    t_clean = normalize_spaces(tema).lower()
    if t_clean in CATALOGO:
        return CATALOGO[t_clean]
    if not t_clean or t_clean == "none":
        return "Sin Definir / Pendiente"
    return "Sin Definir / Pendiente" # Defaults unmatched to Sin Definir

def normalize_modalidad(mod):
    m = normalize_spaces(mod).lower()
    if "presencial" in m: return "Presencial"
    if "virtual" in m: return "Virtual"
    return "Sin dato"

def normalize_prof(prof):
    p = normalize_spaces(prof)
    if not p or p.lower() == "none": return "Sin asignar"
    if p == "Gio": return "Giovanna"
    if p == "Cristian": return "Cristián"
    return p

def extract():
    file_path = 'asesorias_v2.xlsx'
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        print(f"Error loading Excel: {e}")
        sys.exit(1)

    sheet = wb.active
    
    col_nombres = 2
    col_entidad = 18
    
    rondas_cols = [
        {"num": 1, "fecha": 24, "hora": 25, "prof": 26, "mod": 27, "tema": 28, "horas": 29},
        {"num": 2, "fecha": 30, "hora": 31, "prof": 32, "mod": 33, "tema": 34, "horas": 35},
        {"num": 3, "fecha": 36, "hora": 37, "prof": 38, "mod": 39, "tema": 40, "horas": 41},
        {"num": 4, "fecha": 42, "hora": 43, "prof": 44, "mod": 45, "tema": 46, "horas": 47}
    ]

    data = []
    
    # 36 emprendimientos, starts at row 2
    for r in range(2, sheet.max_row + 1):
        nombre_raw = sheet.cell(row=r, column=col_nombres).value
        if not nombre_raw: continue
        nombre = get_str(nombre_raw)
        
        entidad = get_str(sheet.cell(row=r, column=col_entidad).value)
        if not entidad: continue
        
        emprendimiento = f"{entidad} - {nombre}"
        
        asesorias = []
        for ronda in rondas_cols:
            fecha = get_date(sheet.cell(row=r, column=ronda['fecha']).value)
            
            hora_val = sheet.cell(row=r, column=ronda['hora']).value
            if isinstance(hora_val, datetime.time):
                hora_str = hora_val.strftime("%H:%M")
            else:
                hora_str = get_str(hora_val)
                
            horas_val = sheet.cell(row=r, column=ronda['horas']).value
            horas = float(horas_val) if horas_val is not None and str(horas_val).strip() else None
            
            prof = normalize_prof(get_str(sheet.cell(row=r, column=ronda['prof']).value))
            mod = normalize_modalidad(get_str(sheet.cell(row=r, column=ronda['mod']).value))
            tema = normalize_tema(get_str(sheet.cell(row=r, column=ronda['tema']).value))
            
            # If completely empty, skip
            if fecha or (prof != "Sin asignar") or (tema != "Sin Definir / Pendiente") or (horas is not None):
                asesorias.append({
                    "ronda": ronda['num'],
                    "fecha": fecha,
                    "hora_dia": hora_str,
                    "horas": horas,
                    "profesional": prof,
                    "modalidad": mod,
                    "tema": tema
                })
                
        data.append({
            "id": r,
            "emprendimiento": entidad,
            "emprendedor": nombre,
            "nombre_completo": emprendimiento,
            "asesorias": asesorias
        })
        
    with open('src/data/control_gestion_asesorias.json', 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

if __name__ == '__main__':
    extract()
