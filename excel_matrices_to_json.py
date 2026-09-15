import json
import os
import sys
import openpyxl

def clean_str(val):
    if val is None: return ""
    return str(val).strip()

def is_empty_taller(val):
    val_str = clean_str(val).lower()
    return not val_str or val_str == 'none' or val_str == ''

def get_taller_status(val):
    if is_empty_taller(val):
        return "Pendiente"
    # In Excel, typically "Sí", "No", or boolean. 
    # Just return the clean string to be handled by React, but standardize it.
    v = clean_str(val).lower()
    if 'si' in v or 's' == v or 'sí' in v:
        return 'Sí'
    if 'no' in v:
        return 'No'
    return clean_str(val)

def process_matrices(excel_path, out_json_path):
    print(f"Leyendo {excel_path}...")
    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
    except Exception as e:
        print(f"Error cargando Excel: {e}")
        sys.exit(1)
        
    data = {
        "emprendedores": [],
        "jovenes": []
    }
    
    # 1. Emprendedores (Matriz seguimiento)
    sheet_emp = wb['Matriz seguimiento']
    for row in range(10, 40): # rows 10 to 39
        # Nº (1), Nombre (2), Emprendimiento (3), Diagnostico(4), Taller 1(5), T2(6), T3(7), T4(8), T5(9), T6(10), % (11), Horas (12)
        idx = clean_str(sheet_emp.cell(row=row, column=1).value)
        nombre = clean_str(sheet_emp.cell(row=row, column=2).value)
        if not nombre or nombre == 'None':
            continue
            
        emprendimiento = clean_str(sheet_emp.cell(row=row, column=3).value)
        diag = get_taller_status(sheet_emp.cell(row=row, column=4).value)
        t1 = get_taller_status(sheet_emp.cell(row=row, column=5).value)
        t2 = get_taller_status(sheet_emp.cell(row=row, column=6).value)
        t3 = get_taller_status(sheet_emp.cell(row=row, column=7).value)
        t4 = get_taller_status(sheet_emp.cell(row=row, column=8).value)
        t5 = get_taller_status(sheet_emp.cell(row=row, column=9).value)
        t6 = get_taller_status(sheet_emp.cell(row=row, column=10).value)
        
        perc_val = sheet_emp.cell(row=row, column=11).value
        try:
            perc = float(perc_val)
        except:
            perc = 0.0
            
        horas_val = sheet_emp.cell(row=row, column=12).value
        try:
            horas = float(horas_val)
        except:
            horas = 0.0
            
        data["emprendedores"].append({
            "Nº": idx,
            "Nombre del emprendedor": nombre,
            "Emprendimiento": emprendimiento,
            "Diagnóstico inicial": diag,
            "Taller 1": t1,
            "Taller 2": t2,
            "Taller 3": t3,
            "Taller 4": t4,
            "Taller 5": t5,
            "Taller 6": t6,
            "% asistencia talleres": perc,
            "Horas acompañamiento realizadas": horas
        })

    # 2. Estudiantes (Matriz Puerto Serviez)
    sheet_est = wb['Matriz Puerto Serviez']
    for row in range(10, 30): # rows 10 to 29
        # Nº (1), Nombre (2), Idea (3), T1 (4), T2 (5), T3 (6), % (7), Horas (8)
        idx = clean_str(sheet_est.cell(row=row, column=1).value)
        nombre = clean_str(sheet_est.cell(row=row, column=2).value)
        if not nombre or nombre == 'None':
            continue
            
        idea = clean_str(sheet_est.cell(row=row, column=3).value)
        t1 = get_taller_status(sheet_est.cell(row=row, column=4).value)
        t2 = get_taller_status(sheet_est.cell(row=row, column=5).value)
        t3 = get_taller_status(sheet_est.cell(row=row, column=6).value)
        
        perc_val = sheet_est.cell(row=row, column=7).value
        try:
            perc = float(perc_val)
        except:
            perc = 0.0
            
        horas_val = sheet_est.cell(row=row, column=8).value
        try:
            horas = float(horas_val)
        except:
            horas = 0.0
            
        data["jovenes"].append({
            "Nº": idx,
            "Nombre del joven": nombre,
            "Idea / modelo de negocio": idea,
            "Taller 1: Mentalidad emprendedora": t1,
            "Taller 2: Identificación de oportunidades": t2,
            "Taller 3: Modelo de negocio Canvas": t3,
            "% asistencia talleres": perc,
            "Horas acompañamiento realizadas": horas
        })

    with open(out_json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        
    print("Éxito. Matriz procesada.")

if __name__ == '__main__':
    process_matrices('C:\\Users\\jeo20\\Desktop\\Matrices de seguimiento con fechas y estados para programa.xlsx', 'src/data/control_gestion_matrices.json')
